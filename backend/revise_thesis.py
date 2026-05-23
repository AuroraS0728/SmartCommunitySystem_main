from __future__ import annotations

import os
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE_PATH = Path(r"C:\Users\17552\Desktop\论文资料_2026-05-19\宋庆涵.docx")
OUTPUT_PATH = Path(r"C:\Users\17552\Desktop\论文资料_2026-05-19\宋庆涵_定向修改版_工程实现向_2026-05-22.docx")


def set_paragraph_text(paragraph, text: str) -> None:
    paragraph.text = text


def set_cell_text(cell, text: str) -> None:
    cell.text = ""
    lines = text.split("\n")
    p = cell.paragraphs[0]
    p.text = lines[0] if lines else ""
    for line in lines[1:]:
        run = p.add_run()
        run.add_break(WD_BREAK.LINE)
        run.add_text(line)


def insert_paragraph_after(paragraph, text: str = "", style: str | None = None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.text = text
    return new_para


def insert_table_after(doc: Document, paragraph, rows: list[list[str]], style: str = "Table Grid"):
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = style
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            set_cell_text(table.cell(r_idx, c_idx), value)
    paragraph._p.addnext(table._tbl)
    spacer = OxmlElement("w:p")
    table._tbl.addnext(spacer)
    spacer_para = Paragraph(spacer, paragraph._parent)
    return table, spacer_para


def main() -> None:
    if not SOURCE_PATH.exists():
        raise FileNotFoundError(SOURCE_PATH)
    shutil.copyfile(SOURCE_PATH, OUTPUT_PATH)
    doc = Document(str(OUTPUT_PATH))

    paras = doc.paragraphs

    # Global wording cleanup for image audit naming.
    for p in paras:
        if not p.text:
            continue
        text = p.text
        text = text.replace("二手交易图片审核", "二手交易图片真实性辅助审核")
        text = text.replace("Second-hand Trading Image Review", "Second-hand Trading Image Authenticity Assisted Review")
        p.text = text

    # Abstract.
    set_paragraph_text(paras[63], "随着智慧社区建设持续推进，中小型社区的物业服务正由线下分散处理逐步转向多端协同和数据驱动管理。但在实际运行中，仍普遍存在报修派工依赖人工经验、工单时效缺乏连续跟踪、费用催缴方式单一、上门服务安全核验不足以及社区互动信息可信度不高等问题。针对上述问题，本文围绕中小型社区物业业务场景，设计并实现了基于 Spring Boot 的中小型智慧社区物业服务管理系统。")
    set_paragraph_text(paras[64], "系统采用前后端分离架构。后端基于 Spring Boot 构建业务接口，使用 MyBatis-Plus 完成数据访问，以 MySQL 存储用户、房产、报修工单、费用账单、访客邀请、二手交易和推荐规则等核心数据，并使用 Redis 保存验证码、访客动态令牌和工单上门核验等时效性信息。前端包括 Vue 3 开发的 Web 管理端，以及面向业主、家政维修人员和安防人员的微信小程序端。")
    set_paragraph_text(paras[65], "在功能实现方面，系统完成了在线报修、智能派单、工单优先级计算、SLA 时效监控、账单缴费、逾期提醒、访客核验、个性化推荐和二手交易图片真实性辅助审核等业务。智能派单、优先级判断、推荐和催缴等能力主要采用轻量级规则实现；图片真实性辅助审核通过模型服务接口返回真实图片概率和伪造图片概率，再由系统生成风险等级和审核状态。上述结果主要用于为物业管理员提供辅助判断，不直接替代人工决策。系统测试表明，主要业务流程能够正常运行，可为中小型社区物业服务的信息化建设提供工程实现参考。")
    set_paragraph_text(paras[66], "关键词：智慧社区；物业服务；Spring Boot；微信小程序；智能派单")

    set_paragraph_text(paras[79], "With the continuous development of smart communities, property services in small and medium-sized neighborhoods are shifting from fragmented offline handling to coordinated digital management. In practice, however, repair dispatching still relies heavily on manual experience, work-order timeliness is difficult to track, payment reminders are handled in a single way, on-site verification is weak, and the credibility of community trading information is uneven. To address these issues, this thesis designs and implements a smart community property service management system based on Spring Boot.")
    set_paragraph_text(paras[80], "The system adopts a front-end and back-end separated architecture. The back end is built with Spring Boot, uses MyBatis-Plus for data access, stores core business data such as users, properties, repair orders, fee bills, visitor invitations, second-hand trading records and recommendation rules in MySQL, and uses Redis to manage verification codes, temporary visitor tokens and other time-sensitive data. The front end consists of a Vue 3 web administration panel and WeChat Mini Program clients for owners, housekeeping maintenance workers and security staff.")
    set_paragraph_text(paras[81], "In terms of implementation, the system supports online repair requests, intelligent dispatching, work-order priority grading, SLA monitoring, fee payment, overdue reminders, visitor verification, personalized recommendation and assisted authenticity review for second-hand trading images. Lightweight rules are used for dispatching, priority grading, recommendation and reminder generation, while image review is completed through an external model service that returns the probabilities of real and fake images. These intelligent functions provide decision support for property administrators rather than replacing manual judgment. The test results show that the main business processes can run normally and the system can serve as an engineering reference for property service informatization in small and medium-sized communities.")
    set_paragraph_text(paras[82], "Key words: Smart community; Property service; Spring Boot; WeChat Mini Program; Intelligent dispatching")

    # 1.3.1
    set_paragraph_text(paras[196], "本系统面向中小型社区物业服务场景，围绕多角色协同、报修流转、费用处理和社区互动等核心业务展开设计。结合论文研究目标与现有源码实现，本文主要研究内容如下。")
    research_items = {
        197: "设计并实现多角色协同的物业服务基础功能。系统围绕业主、物业管理员、家政维修人员和安防人员划分业务入口与权限范围，使报修、缴费、访客核验、投诉反馈和社区互动能够在统一后端服务上协同运行。",
        198: "设计在线报修与智能派单功能。系统支持业主提交报修类型、预约日期、预约时段、问题描述和图片信息，并根据服务类别、专业标签、当前任务量和可服务状态为物业管理员提供派单参考，管理员仍可手动调整派工结果。",
        199: "设计上门服务核验机制。系统提供 SeetaFace6.0 活体人脸核验能力和现场验证码核验能力，用于记录家政维修人员上门核验结果，并将核验结果写入工单参与人员记录，为后续工单状态流转提供依据。",
        200: "设计工单优先级判定与时效监控功能。系统根据报修描述中的紧急关键词和业主信用情况生成工单优先级，并在工单创建、派单和处理中设置时限节点，对超时工单生成提醒与催办记录。",
        201: "设计费用账单与逾期提醒功能。系统支持物业费等账单生成、查询、积分支付和模拟支付，并对逾期账单进行后台扫描，根据欠费金额和逾期时长生成消息提醒或物业待办。",
        202: "设计社区互动与二手交易图片真实性辅助审核功能。系统支持二手交易、失物招领等社区互动业务，并在二手交易发布流程中调用图片检测服务，依据伪造图片概率生成风险等级和审核状态，为物业管理员提供辅助判断。",
        203: "设计个性化推荐与积分信用记录功能。系统根据业主显式画像、隐式关键词和历史水电类报修次数匹配推荐规则，并分别维护积分消费记录和信用分变动日志，便于后续扩展更多行为评价规则。",
        204: "完成系统实现与测试验证。系统在后端接口、数据库结构和多端页面基础上完成联调，并结合白盒测试、黑盒测试和接口测试对核心业务流程进行验证。"
    }
    for idx, text in research_items.items():
        set_paragraph_text(paras[idx], text)
    set_paragraph_text(paras[205], "上述研究内容不再停留于功能罗列，而是重点关注业务规则如何通过控制层、业务服务层、数据访问层、数据表结构和测试用例落地，从而增强论文的工程实现性。")

    # 1.3.2
    set_paragraph_text(paras[207], "为完成系统设计与实现，本文按照需求分析、系统设计、功能实现和测试验证的顺序开展研究。系统采用前后端分离架构，Spring Boot 负责统一提供业务接口，Controller 层负责接收请求和基础校验，Service 层负责业务规则处理，MyBatis-Plus 负责实体映射与数据访问，MySQL 用于保存核心业务数据。")
    set_paragraph_text(paras[208], "在支撑组件方面，Redis 用于保存登录相关时效数据、访客动态令牌、工单上门验证码及核验状态等缓存信息；Vue 3 用于实现物业 Web 管理端；微信小程序用于实现业主端、家政维修端和安防端。该技术组合能够同时满足多端访问、统一接口复用和后续功能扩展的需要。")
    set_paragraph_text(paras[209], "在关键功能实现思路上，智能派单、工单分级、SLA 监控、逾期提醒和个性化推荐等能力主要通过轻量级规则实现。系统在控制层完成参数校验，在业务服务层完成规则计算和状态流转，在数据访问层完成工单、账单、提醒记录和推荐规则的持久化，从而形成可追踪的工程实现路径。")
    set_paragraph_text(paras[210], "对于二手交易图片真实性辅助审核，系统通过 HTTP 接口调用外部图片检测服务，模型返回真实图片概率和伪造图片概率，系统再根据阈值生成风险等级和审核状态，并决定是否进入人工复核流程。该结果用于辅助物业管理员审核，不直接替代人工决策。")

    # 2.1.3 and 2.2.1 requirement wording cleanup.
    set_paragraph_text(paras[233], "根据系统业务范围，本文将物业服务管理系统划分为用户与房产管理、报修工单管理、智能派单与时效监控、费用账单管理、逾期提醒管理、公告活动管理、访客通行管理、投诉评价管理、邻里互动管理、积分与信用记录管理、个性化推荐管理和数据统计分析等功能模块。")
    set_paragraph_text(paras[234], "其中，用户与房产管理模块用于完成业主、物业管理员、家政维修人员和安防人员的账号管理、身份认证和权限控制；报修工单管理模块用于支持报修提交、工单查询、派工处理、维修反馈和服务评价；智能派单与时效监控模块用于根据服务类别、专业标签、当前任务量、可服务状态和时限节点生成派工参考与超时提醒；费用账单管理模块用于完成账单生成、查询、积分支付和模拟支付。")
    set_paragraph_text(paras[235], "访客通行管理模块用于生成访客邀请码、动态二维码和访客核验记录；邻里互动管理模块用于支持二手交易、失物招领等社区互动内容，并在二手交易发布流程中接入图片真实性辅助审核；积分与信用记录模块用于保存积分充值、积分消费和信用分变动日志；个性化推荐模块用于根据画像字段、历史行为关键词和推荐规则返回推荐结果。")
    set_paragraph_text(paras[236], "通过上述功能模块划分，系统能够较为完整地覆盖中小型社区物业服务的主要业务需求，并为后续工程实现和维护扩展提供较清晰的职责边界。")

    set_paragraph_text(paras[243], "（1）用户与房产管理需求。系统应支持不同角色用户的登录、身份认证和权限控制。业主首次使用系统时，可以通过房产绑定建立个人账号与房屋之间的关联，并查看个人资料、名下房产、积分余额和信用分等信息。物业管理员需要对楼栋、单元、房屋、入住状态以及业主关联关系进行维护，为后续报修、缴费、访客通行等业务提供基础数据。")
    set_paragraph_text(paras[244], "（2）报修工单管理需求。系统应支持业主在线提交报修申请，提交内容包括服务类别、预约日期、预约时段、问题描述和现场图片等。报修申请提交后，系统生成工单并向业主展示处理进度。物业管理员可以查看工单详情并完成派工处理；家政维修人员应能够查看分配给自己的工单，完成接单、上门核验、状态更新和维修凭证上传。")
    set_paragraph_text(paras[245], "（3）智能派单与时效监控需求。系统应结合服务类别、问题描述、紧急关键词、家政维修人员专业标签、当前任务数量和可服务状态，为物业管理员提供派工参考。对于超过处理时限的工单，系统应生成超时提醒，并记录催办信息和逾期次数，为物业管理员提供辅助判断。")
    set_paragraph_text(paras[246], "（4）费用账单与逾期提醒需求。系统应支持物业费、停车费等账单的生成、查询和状态维护。业主可以查看账期、应缴金额、截止日期和缴费状态，并通过系统提供的积分支付或模拟支付方式完成缴费。对于逾期账单，系统应根据欠费金额和逾期天数生成消息提醒或物业待办，避免重复生成同类提醒。")
    set_paragraph_text(paras[247], "（5）访客通行管理需求。系统应支持业主填写访客姓名、手机号、来访时间等信息，并生成访客邀请码、二维码或动态通行凭证。安防人员应能够对访客凭证进行核验，记录通行结果，并对凭证过期、黑名单命中、动态令牌失效等异常情况给出提示。")
    set_paragraph_text(paras[248], "（6）投诉评价与反馈处理需求。系统应支持业主提交投诉建议并上传相关图片，物业管理员完成受理、回复和办结处理，业主能够查看回复内容并进行满意度评价。系统可对投诉文本进行初步分析，为物业管理员提供辅助预警信息。")
    set_paragraph_text(paras[249], "（7）邻里互动、积分与信用记录需求。系统应支持二手交易、失物招领和社区互动等功能，并支持积分充值、积分消费和信用分记录。当前源码可确认的信用分变动场景包括报修五星评价、二手交易成功和管理员人工调整；其他跨业务信用规则可作为后续扩展内容，需人工确认。")
    set_paragraph_text(paras[250], "（8）个性化推荐与图片审核需求。系统应支持根据用户显式画像、历史水电类报修次数和隐式关键词返回推荐内容，并支持对二手交易图片进行真实性辅助审核。系统应根据检测结果给出正常、疑似异常或需人工复核等提示，检测结果主要用于辅助物业管理员判断。")
    set_paragraph_text(paras[251], "除上述基础业务功能外，系统还需要提供一定的智能辅助能力。例如，在报修业务中，系统可根据工单内容和家政维修人员状态进行优先级判断和人员推荐；在费用业务中，系统可根据逾期时间和账单金额生成不同提醒方式；在推荐业务中，系统可根据用户画像和规则配置展示相应内容；在邻里互动业务中，系统可对二手交易图片进行真实性辅助识别。上述功能均用于提供辅助参考，不直接替代人工判断。")

    # 2.3.3 intro + image audit caption.
    set_paragraph_text(paras[279], "本文结合智慧社区物业服务的主要业务流程，对系统核心功能进行用例分析。用例描述从参与者、前置条件、后置条件、基本流程和可选流程等方面展开，用于明确房产绑定、在线报修、维修处理、账单缴费、访客核验、个性化推荐和二手交易图片真实性辅助审核等功能的业务边界，并为后续系统设计提供依据。")
    set_paragraph_text(paras[298], "表2-9 二手交易图片真实性辅助审核用例描述")
    set_paragraph_text(paras[299], "Tab.2-9 Use Case Description of Assisted Authenticity Review for Second-hand Trading Images")

    # 2.4
    set_paragraph_text(paras[305], "在邻里互动业务中，二手交易信息通过 second_hand 对象保存商品标题、价格、联系方式和图片地址列表。商品图片在发布流程中经过真实性辅助检测后，可生成对应的图片审核结果对象，用于记录真实图片概率、伪造图片概率、风险等级和审核状态，为后续人工复核提供依据。")

    # 3.x wording cleanup for unsupported claims.
    set_paragraph_text(paras[316], "在线报修与智能派单业务主要用于完成业主报修、工单生成、物业派单和家政维修人员接单等操作。业主登录微信小程序后，填写报修类型、预约日期、预约时段、问题描述，并上传现场图片。前端将报修信息提交至后端后，RepairController 先对报修信息完整性和预约时段可用性进行校验；校验通过后，系统调用 RepairGradingService 计算工单优先级，再结合 WorkerRecommendService 与工单候选人评分逻辑生成推荐人员，并将工单信息写入 repair_order。物业管理员随后可在管理端查看推荐结果并完成派工处理。")
    set_paragraph_text(paras[334], "在线报修协作主要体现业主小程序、后端控制器、业务服务、数据访问对象、物业管理端和家政维修端之间的职责分工。业主小程序负责采集报修类型、预约时间、问题描述和图片等信息，并向后端提交报修请求。RepairController 作为接口入口，负责接收请求并进行参数校验；RepairGradingService 负责计算工单优先级；候选家政维修人员的筛选与评分逻辑由 RepairController 中的工单评分过程及 WorkerRecommendService 共同完成；RepairOrderMapper 与 RepairOrderWorkerMapper 负责工单和参与人员数据读写。")
    set_paragraph_text(paras[413], "报修工单的状态从业主提交报修申请开始，初始状态为“待派单”。物业管理员确认派单，或系统根据服务类别、专业标签、当前任务量和可服务状态给出推荐人员后，工单进入“处理中”状态。家政维修人员上门前可进行人脸活体验证或现场验证码核验，核验结果写入工单参与人员记录；核验通过后开始维修处理。维修完成后，家政维修人员上传维修凭证并填写费用明细，工单进入“待评价”状态。业主确认维修结果并完成评价后，工单状态变为“已完成”；若业主取消报修或业务处理异常，工单可进入“已取消”状态。")

    # 4.4 section text.
    set_paragraph_text(paras[457], "详细类设计是在系统架构设计的基础上，对核心业务模块中的主要实现类及其关系进行进一步说明。本系统采用 Controller、Service、Mapper 和 Entity 分层设计方式。结合现有源码，4.4 节不再停留于模块功能罗列，而是重点说明智能派单、工单优先级、积分与信用记录、逾期提醒、个性化推荐、SLA 监控和二手交易图片真实性辅助审核等功能如何通过具体类、接口、数据表和测试用例落地。")
    set_paragraph_text(paras[462], "费用管理模块以 FeeController 作为账单查询、积分支付和模拟支付入口，以 FeeBillingService 负责账单生成，以 PaymentReminderService 负责逾期提醒扫描和提醒落库。FeeBillMapper、PaymentReminderMapper 和 PropertyTaskMapper 分别负责账单、提醒记录和物业待办的数据访问，形成从账单生成、状态校验到提醒分流的闭环。")
    set_paragraph_text(paras[466], "访客安防模块以 AccessController 作为接口入口，主要负责访客邀请码生成、动态令牌生成、访客凭证核验、通行通知和黑名单处理。VisitorInviteMapper、AccessTokenMapper、VisitorBlacklistMapper 和 VisitorNotifyMapper 分别负责访客邀请、临时令牌、黑名单和通知记录的数据访问。该模块通过凭证有效期、使用次数、动态令牌状态和黑名单校验等机制实现安防核验流程。")
    set_paragraph_text(paras[470], "个性化推荐模块以 RecommendController 作为接口入口，主要提供首页推荐和弹窗推荐接口。RecommendRuleService 根据用户显式画像、历史水电类报修次数和隐式关键词执行规则匹配；ImplicitProfileService 负责从二手交易标题和失物招领描述中提取关键词并写入用户画像表；RecommendRuleMapper 和 UserImplicitProfileMapper 分别负责推荐规则和隐式画像数据访问。")
    set_paragraph_text(paras[474], "二手交易图片真实性辅助审核模块主要由 NeighborController、NeighborServiceImpl、ImageAuditService、ImageAuditClient 和 ImageAuditResultMapper 共同完成。业务服务层负责解析图片地址、调用模型服务、计算风险等级、保存审核结果并在必要时将交易状态改为待审核；检测结果只作为辅助参考，不直接替代物业管理员的人工审核。")

    # 4.5 database wording cleanup.
    set_paragraph_text(paras[483], "为支持系统中的智能辅助功能，数据库设计中增加了信用分日志、逾期提醒记录、物业待办、推荐规则、用户隐式画像和图片审核结果等数据结构。同时，系统还通过 user.points、points_recharge_record 和 points_consumption_record 维护积分余额与积分消费过程。上述数据共同为智能派单、逾期提醒、服务推荐和二手交易图片真实性辅助审核提供数据基础。")
    set_paragraph_text(paras[496], "在二手交易图片真实性辅助审核业务中，second_hand 表通过 images 字段保存同一商品的一张或多张图片地址，image_audit_result 表按图片维度保存检测结果。该结果表记录真实图片概率、伪造图片概率、阈值、风险等级和审核状态，并可在商品详情查询时汇总为整体审核状态，为后续人工复核和交易信息管理提供数据依据。")

    # 5.3
    set_paragraph_text(paras[587], "本节选取系统中具有代表性的核心功能模块进行实现说明。除运行界面外，重点补充前后端交互、控制层处理、业务服务层规则计算和数据落库过程，以体现系统的工程实现路径。")
    set_paragraph_text(paras[589], "费用缴纳与停车缴费模块主要用于业主查看物业费、停车费等账单信息，并支持积分支付、模拟微信支付和缴费状态查看。前端请求账单列表后，后端通过 FeeController 查询名下房产对应账单；业主选择积分支付时，系统调用 PointsService 校验积分余额并更新 fee_bill、points_consumption_record 等数据；选择模拟微信支付时，系统通过 /api/fee/pay/wechat 生成模拟支付参数，并在回调接口中更新账单状态。")
    set_paragraph_text(paras[593], "报修提交与工单详情模块用于业主提交维修需求并查看工单处理进度。业主端提交 serviceMajor、serviceSubType、category、appointmentDate、appointmentTimeSlot、description 和 images 后，RepairController 先校验必填项、预约日期格式、预约时段有效性以及可预约时段剩余人手，再调用 RepairGradingService 计算优先级，写入 repair_order.priority、suggested_worker_id 和 sla_deadline，最后保存工单。详情页再通过 repair_order_worker、repair_fee_detail、repair_fee_bill 等关联数据展示处理进度、参与人员、核验状态和费用信息。")
    set_paragraph_text(paras[597], "活动专区与首页推荐页面主要用于展示社区活动和服务推荐内容。当前端请求推荐内容时，RecommendController 调用 RecommendRuleService，系统读取当前业主显式画像、历史水电类报修次数以及 user_implicit_profile 中的关键词，再与 recommend_rule 中启用规则进行匹配。首页普通推荐默认返回不超过 3 条内容，弹窗推荐默认返回不超过 2 条内容。")
    set_paragraph_text(paras[601], "二手交易模块用于社区闲置物品发布与浏览。业主发布商品时，前端提交标题、类别、价格、描述、联系方式和 images 字段；NeighborController 接收请求后交由 NeighborServiceImpl 保存 second_hand 记录，并调用 ImageAuditService 对每张图片执行真实性辅助审核。系统把审核结果写入 image_audit_result；若任一图片达到人工复核条件，则把商品状态调整为待审核。")
    set_paragraph_text(paras[605], "在二手交易发布流程中，图片真实性辅助审核采用“商品发布接口触发 + 模型服务返回结果落库”的实现方式。ImageAuditClient 通过 HTTP Multipart 请求调用模型服务接口，接收真实图片概率、伪造图片概率和阈值信息；ImageAuditService 再根据 fake-threshold 和 review-threshold 生成风险等级与审核状态，并把结果返回给前端展示。该功能只提供辅助判断，不直接替代物业管理员审核。")
    set_paragraph_text(paras[610], "访客邀请与核验模块用于生成和管理访客通行凭证。业主端调用 AccessController 生成邀请码或动态二维码；安防端通过 /api/access/verify-invite 或 /api/access/verify-token 提交核验请求，后端据此校验有效期、使用次数、黑名单状态以及动态令牌是否仍为最新令牌，核验通过后保存 visitor_invite 使用信息并生成业主通知。")
    set_paragraph_text(paras[615], "家政维修端主要面向家政维修人员使用，支持工单查看、消息提醒和工单详情处理等功能。维修人员登录后可通过 WorkerController.tasks 查询分配给自己的工单；后端会结合 assignee 字段和 repair_order_worker 参与人员记录限制数据范围，使家政维修人员只查看与自己相关的工单。")
    set_paragraph_text(paras[619], "上门核验与维修处理页面主要用于到场后的身份核验和维修结果提交。系统提供 WorkerController.verify-code 的现场验证码核验接口，以及 FaceController.liveVerifyMulti 的人脸活体验证接口，两类核验结果都会写入 repair_order_worker.verify_passed 和 verify_pass_time。源码未发现必须同时通过两类核验的强制逻辑，该组合策略需人工确认。家政维修人员完成核验后，可在 RepairController.updateStatus 中上传维修前后图片、填写技术费和材料费等明细，并把结果汇总到 repair_fee_detail 和 repair_fee_bill。")
    set_paragraph_text(paras[624], "物业安防端主要面向安防人员和物业管理员使用。安防人员可在工作台进入门禁核验页面，对访客二维码、邀请码或动态令牌进行验证；物业管理员可通过 RepairController 的费用异议接口查看业主提交的异议申请并完成审核，也可通过 ComplaintController 查看、处理和回复投诉建议，从而形成访客核验、费用争议处理和投诉反馈的线上协同流程。")

    # 6.2 and 6.3
    set_paragraph_text(paras[638], "由于在线报修与智能派单模块同时涉及 RepairController.submit、RepairGradingService.calculatePriority、候选家政维修人员筛选逻辑以及 SlaMonitorService.scanAndUrge 等关键处理环节，且这些环节直接影响工单生成、优先级判定、派工建议和超时催办结果，因此本文仍选取该模块作为白盒测试对象。测试重点放在输入校验、优先级分支、候选人筛选分支和超时扫描分支是否能够按照设计执行。")
    set_paragraph_text(paras[640], "在线报修与智能派单模块主要用于完成业主报修申请提交、系统生成工单、计算工单优先级、推荐家政维修人员以及物业管理员确认派单等操作。业主在小程序端填写服务类型、预约日期、预约时段、问题描述并上传图片后，系统首先对输入信息进行完整性与格式校验；若信息不完整或预约时段无效，则立即返回错误提示；若校验通过，则继续判断报修内容中是否包含紧急关键词，并结合业主信用情况生成工单优先级。")
    set_paragraph_text(paras[641], "在完成工单优先级计算后，系统根据服务类别、专业标签、证书信息、当前任务量和可服务状态筛选候选家政维修人员。若存在候选人，则在保存工单时写入 suggested_worker_id；若预约时段没有可用人员，则在提交阶段直接返回失败提示。对于已经生成并进入待派单或处理中状态的工单，SlaMonitorService 还会基于 sla_deadline 执行超时扫描，生成 sys_message 和 repair_urge_log 记录。")
    set_paragraph_text(paras[652], "除控制流图对应的 5 条基础独立路径外，本文补充 1 条针对 SLA 扫描的派生测试路径，用于验证超时催办逻辑是否能够正确生成提醒记录。")
    set_paragraph_text(paras[658], "由表6-3可知，在线报修与智能派单模块在不同输入条件下能够覆盖报修信息校验、优先级计算、候选人筛选和超时催办等主要执行路径。测试结果表明，该模块在 priority、suggested_worker_id、sla_deadline、repair_urge_log 和 sys_message 等关键数据处理上能够按照预期执行。")

    set_paragraph_text(paras[660], "黑盒测试主要从用户使用角度验证系统功能是否符合需求。测试过程中不关注程序内部实现，而是通过输入测试数据、观察接口返回、状态变化和数据落库结果，检查系统输出是否符合预期。本系统黑盒测试主要包括功能测试、等价类划分测试、边界值测试、状态转换测试、场景测试和接口测试，并额外关注积分与信用记录、逾期提醒、个性化推荐和图片审核等智能辅助功能。")
    set_paragraph_text(paras[662], "功能测试用于验证系统各功能模块是否能够按照需求正常运行。本文根据系统功能需求，对房产绑定、在线报修、智能派单、维修处理、费用缴纳、访客核验、投诉评价、个性化推荐、逾期提醒、积分支付和二手交易图片真实性辅助审核等功能进行测试。")
    set_paragraph_text(paras[667], "等价类划分法用于将输入数据划分为有效等价类和无效等价类，从而减少测试用例数量。本系统中，登录、房产绑定、报修提交、账单支付、访客凭证核验和二手交易发布等功能均涉及输入数据校验，因此适合采用等价类划分法进行测试。")
    set_paragraph_text(paras[675], "边界值测试用于验证系统在临界值附近是否能够正确处理。本系统中报修描述长度、账单金额、逾期天数、访客动态令牌有效期和图片地址列表等数据均存在边界限制，因此需要进行边界值测试。")
    set_paragraph_text(paras[681], "状态转换测试适用于状态变化较为明显的业务对象。本系统中报修工单、费用账单、访客凭证、投诉建议和二手交易审核结果均具有明确的状态流转过程，因此选取这些对象进行状态转换测试。")
    set_paragraph_text(paras[686], "场景测试根据用例描述中的基本事件流和可选事件流设计测试场景，适合验证完整业务流程是否能够顺利执行。本文选取在线报修完整流程、费用缴纳与逾期提醒流程、访客邀请与安防核验流程、个性化推荐流程和二手交易发布审核流程等典型业务场景进行测试。")
    set_paragraph_text(paras[691], "本系统采用前后端分离架构，前端通过接口与后端进行数据交互，因此需要对主要接口进行测试。接口测试主要验证请求参数、返回结果、权限控制和异常提示是否符合设计要求，并重点检查报修提交、访客核验、推荐返回、积分支付和二手交易发布流程中的审核结果。")
    set_paragraph_text(paras[696], "接口测试结果表明，系统主要接口能够正确处理正常请求和异常请求，权限控制、参数校验、状态更新和图片审核结果回写基本符合预期。")

    # 6.3.7
    set_paragraph_text(paras[698], "由于二手交易图片真实性辅助审核功能主要依赖 Vision Transformer 模型对图片进行二分类判断，其测试重点不在于传统业务代码路径覆盖，而在于模型输出结果与系统审核状态是否一致。因此，本文将该功能放在黑盒测试和模型效果测试中进行验证，主要从数据划分、模型评价指标、接口返回结果和系统审核状态四个方面展开测试。")
    set_paragraph_text(paras[699], "当前模型采用 real/fake 二分类结构，输出类别为 real 和 fake，并返回真实图片概率与伪造图片概率。系统根据伪造图片概率和设定阈值生成风险等级与审核状态，例如正常、疑似异常或需人工复核。该检测结果主要作为物业管理员审核参考，不直接作为商品发布与否的唯一依据。")
    set_paragraph_text(paras[701], "本文构建了面向二手交易图片真实性检测的二分类数据集，类别包括真实图片和伪造图片。真实图片用于模拟正常二手商品或生活场景图片，伪造图片样本则来自非真实图片集合。为了增强模型对伪造风险图片的识别能力，实验中对部分补充样本进行了拆分和清洗处理，最终形成训练集、验证集和测试集。")
    set_paragraph_text(paras[711], "从测试结果可以看出，模型对真实图片和伪造图片具有一定区分能力，能够为系统发布流程提供辅助风险提示。需要说明的是，当前模型属于 real/fake 二分类模型，不能直接判定伪造图片的具体来源或生成方式。因此，系统仅将模型输出转换为风险等级和审核状态，并把高风险结果转入人工复核流程。")

    # 6.4-6.6
    set_paragraph_text(paras[713], "经过白盒测试和黑盒测试，系统主要功能能够满足需求分析阶段提出的基本要求。测试结果表明，系统在正常输入、异常输入、边界条件和完整业务场景下均能够完成相应处理，并能够给出合理提示。下面从软件能力、缺陷和限制、建议以及测试结论四个方面进行说明。")
    set_paragraph_text(paras[715], "通过测试可以看出，本系统已具备较完整的物业服务业务处理能力。业主端能够完成房产绑定、在线报修、费用查询、积分支付、访客邀请、投诉建议、二手交易发布和推荐内容查看等操作；物业管理端能够完成房产档案维护、工单派发、费用管理、逾期提醒生成、投诉回复和统计查看；家政维修端能够完成工单接收、上门核验、维修凭证上传和状态更新；安防端能够完成访客凭证核验和通行记录保存。")
    set_paragraph_text(paras[716], "在核心业务流程方面，系统能够支持报修申请、工单分级、推荐派工、维修反馈和业主评价的完整闭环；费用管理功能能够支持账单生成、账单查询、积分支付、模拟支付和逾期提醒；访客管理功能能够支持访客凭证生成、动态令牌核验和通行记录保存；个性化推荐功能能够根据用户画像和规则配置展示对应内容。整体来看，系统能够满足中小型社区物业服务的基本使用需求。")
    set_paragraph_text(paras[718], "虽然系统主要功能测试结果基本符合预期，但仍存在一定不足。首先，费用缴纳功能目前采用模拟支付方式，尚未接入真实第三方支付平台。其次，个性化推荐采用“用户画像 + 规则匹配”的轻量级方式实现，推荐精细度仍有提升空间。再次，现有信用分规则在源码中仅能确认报修五星评价、二手交易成功、人工调整和时间衰减，更多跨业务信用变动规则仍需人工确认。最后，人脸活体验证和现场验证码虽然都已接入工单流程，但源码未发现必须同时通过两类核验的强制逻辑，该联动策略仍有进一步完善空间。")
    set_paragraph_text(paras[720], "针对上述限制，后续可以从以下几个方面继续完善。首先，可进一步接入微信支付等真实支付接口，并补充支付回调、退款处理和账单对账机制。其次，可在积累更多用户行为数据后，优化推荐规则和信用评价规则。再次，可补充更大规模的并发测试和长期稳定性测试。最后，可继续完善上门核验联动策略、模型服务容错处理和图片审核样本集扩充，提高系统在真实社区环境中的适用性。")
    set_paragraph_text(paras[722], "通过对白盒测试和黑盒测试结果的分析可以看出，本系统主要功能基本实现，核心业务流程能够正常运行。系统已经能够支持业主端、物业管理端、家政维修端和安防端之间的信息交互，并在报修派工、逾期提醒、积分支付、个性化推荐和图片审核等环节体现出一定的工程实现能力。")
    set_paragraph_text(paras[723], "测试过程中未发现影响系统基本运行的严重错误。对于真实支付接口未接入、推荐算法较轻量、信用规则覆盖有限和性能测试规模有限等问题，后续可在系统扩展和实际部署过程中继续完善。总体来看，本系统能够满足中小型社区物业服务管理的基本需求，测试结论为基本通过。")

    # 7 conclusion
    set_paragraph_text(paras[726], "本文围绕中小型社区物业服务场景，完成了基于 Spring Boot 的中小型智慧社区物业服务管理系统的设计与实现。系统以多端协同为主线，围绕业主、物业管理员、家政维修人员和安防人员的实际业务流程，完成了需求分析、系统分析、系统设计、功能实现和测试验证等工作。")
    set_paragraph_text(paras[727], "在系统实现层面，后端采用 Spring Boot、MyBatis-Plus、MySQL 和 Redis 形成统一服务支撑，前端采用 Vue 3 和微信小程序构建多角色入口。系统实现了业主端、物业管理端、家政维修端和安防端之间的业务协作，能够支撑在线报修、工单派发、维修反馈、费用缴纳、访客核验、投诉回复和社区互动等主要业务。")
    set_paragraph_text(paras[728], "在工程实现细节上，系统已将部分智能辅助能力落到具体控制器、服务类、数据表和接口流程中。例如，RepairController、RepairGradingService、WorkerRecommendService 和 SlaMonitorService 共同完成了工单提交、优先级判定、候选人推荐和超时催办；PaymentReminderService 将逾期账单扫描结果分流为消息提醒或物业待办；RecommendRuleService 基于显式画像、隐式关键词和规则表达式返回推荐内容；ImageAuditService 与 ImageAuditClient 则在二手交易发布流程中完成图片真实性辅助审核。")
    set_paragraph_text(paras[729], "需要说明的是，这些智能辅助功能主要通过轻量级规则和模型服务接口实现，重点在于为物业管理员提供辅助判断，而不是替代人工决策。与原有偏重功能列举的写法相比，本文更加突出 Controller、Service、Mapper、数据库表和接口之间的协作关系，从而增强论文的工程实现性。")
    set_paragraph_text(paras[730], "系统测试表明，报修提交、派工处理、费用缴纳、逾期提醒、访客核验、投诉回复、个性化推荐和二手交易图片真实性辅助审核等流程能够基本按照预期运行。白盒测试覆盖了输入校验、优先级计算、候选人筛选和 SLA 超时催办等关键路径；黑盒测试覆盖了正常场景、异常场景和接口调用结果。")
    set_paragraph_text(paras[731], "当然，系统仍存在一定不足。首先，真实支付接口尚未接入，当前费用缴纳以积分支付和模拟支付为主。其次，推荐与信用评价规则仍以轻量级实现为主，覆盖范围和精细度有限。再次，图片审核模型的训练样本规模和场景覆盖仍有局限，模型结果只能作为辅助参考。最后，现有并发测试规模较小，尚未在真实社区环境下进行长周期运行验证。")
    set_paragraph_text(paras[732], "后续工作可从真实支付接入、派工与推荐规则优化、信用规则扩展、图片审核样本扩充、热力图可解释分析以及真实社区部署验证等方面继续完善，使系统更贴近实际物业服务应用。")

    # Use case tables.
    table = doc.tables[1]
    set_cell_text(table.cell(8, 1), "1. 业主进入小程序报修页面。\n2. 业主选择服务类型，填写预约日期、预约时段、问题描述并上传图片。\n3. 系统校验报修信息完整性、预约日期格式和预约时段有效性。\n4. 系统校验所选预约时段是否仍有可用家政维修人员。\n5. 系统生成报修工单。\n6. 系统根据报修内容和业主信用情况计算工单优先级。\n7. 系统生成建议派工人员。\n8. 物业管理员在管理端查看工单并采纳推荐或手动选择人员。\n9. 系统保存派单信息，并将工单状态更新为处理中。")
    set_cell_text(table.cell(9, 1), "1. 若业主提交的信息不完整，系统提示补充报修内容。\n2. 若 appointmentDate 格式错误或 appointmentTimeSlot 非法，系统返回参数错误提示。\n3. 若所选预约时段没有可用家政维修人员，系统拒绝提交并提示重新选择时段。\n4. 若系统未生成建议派工人员，物业管理员可手动指定家政维修人员。\n5. 若工单已完成或取消，系统不允许再次派单。")

    table = doc.tables[3]
    set_cell_text(table.cell(9, 1), "1. 若业主无待缴账单，系统显示暂无待缴费用。\n2. 若积分余额不足，系统提示积分不足。\n3. 若账单已缴费，系统不重复扣费。\n4. 若模拟支付失败或回调失败，系统保持账单未缴状态并提示重新支付。")

    table = doc.tables[4]
    set_cell_text(table.cell(9, 1), "1. 若访客手机号命中黑名单，系统拒绝通行。\n2. 若邀请码或动态令牌过期，系统提示凭证无效。\n3. 若动态令牌不是当前最新令牌，系统提示二维码已刷新。\n4. 若访客手机号与邀请信息不一致，系统拒绝核验。\n5. 若业主删除或撤销访客邀请，系统不再允许该凭证通行。")

    table = doc.tables[5]
    set_cell_text(table.cell(8, 1), "1. 业主进入小程序首页。\n2. 小程序向后端请求推荐服务数据或弹窗推荐数据。\n3. 系统读取业主显式画像和隐式关键词。\n4. 系统读取启用状态的推荐规则。\n5. 系统根据规则表达式和条件参数进行匹配。\n6. 首页普通推荐默认返回不超过 3 条内容，弹窗推荐默认返回不超过 2 条内容。\n7. 小程序展示推荐卡片或弹窗内容。\n8. 业主点击后跳转至对应服务页面。")
    set_cell_text(table.cell(9, 1), "1. 若业主未登录，系统不返回个性化推荐结果。\n2. 若无匹配规则，系统返回空列表。\n3. 若请求弹窗推荐，系统仅返回 popupEnabled 为 true 的规则。\n4. 若同一服务被多条规则同时命中，系统按 serviceId 去重后返回。")

    table = doc.tables[6]
    set_cell_text(table.cell(0, 1), "二手交易图片真实性辅助审核")
    set_cell_text(table.cell(2, 1), "本用例描述业主发布二手交易信息时，系统对上传商品图片进行真实性辅助检测，并根据检测结果提示是否需要人工复核的过程。")
    set_cell_text(table.cell(5, 1), "系统保存商品信息和图片检测结果；若存在高风险结果，则对应商品进入待审核状态。")
    set_cell_text(table.cell(6, 1), "1. 业主填写商品名称、价格、描述和联系方式。\n2. 业主提交商品图片地址列表。\n3. 系统保存二手交易记录。\n4. 系统逐张调用图片检测服务。\n5. 模型返回真实图片概率和伪造图片概率。\n6. 系统根据伪造图片概率生成风险等级与审核状态。\n7. 若存在需人工复核结果，则将商品状态调整为待审核；否则保持正常发布状态。")
    set_cell_text(table.cell(7, 1), "1. 若图片地址无法读取或模型检测失败，系统将审核状态记为需人工复核。\n2. 若图片格式不支持或图片资源异常，系统可拒绝发布或转入人工复核，具体前端限制需人工确认。\n3. 若模型服务不可用，系统仍保存商品信息，但将审核状态标记为需人工复核。")

    # White-box tables.
    table = doc.tables[10]
    extra_row = table.add_row().cells
    for c, text in enumerate(["P6", "S1→S2→S3→S4", "工单超过 SLA 截止时间且 30 分钟内未重复催办，系统写入 sys_message 与 repair_urge_log。"]):
        set_cell_text(extra_row[c], text)
    set_cell_text(table.cell(1, 2), "报修信息不完整，系统返回错误提示。")
    set_cell_text(table.cell(2, 2), "普通报修信息完整，预约时段有效，系统生成普通优先级工单并写入建议派工人员。")
    set_cell_text(table.cell(3, 2), "报修描述包含紧急关键词，系统提高工单优先级。")
    set_cell_text(table.cell(4, 2), "所选预约时段无可用家政维修人员，系统拒绝提交。")
    set_cell_text(table.cell(5, 2), "业主信用分较低，系统将普通工单降级为低优先级后保存。")

    table = doc.tables[11]
    extra_row = table.add_row().cells
    new_rows = [
        ["WT01", "P1", "description 为空，serviceMajor 为空。", "系统提示报修信息不完整。", "通过，返回 service type or description is empty。"],
        ["WT02", "P2", "serviceMajor=水电维修，description=厨房水龙头松动，appointmentDate 和 appointmentTimeSlot 合法。", "系统生成普通优先级工单，并写入 suggestedWorkerId。", "通过，工单保存成功并生成建议派工人员。"],
        ["WT03", "P3", "description 包含“漏水”“停电”等紧急关键词。", "系统提高工单优先级为紧急。", "通过，priority 写入为 1。"],
        ["WT04", "P4", "预约时段内对应服务类别无可用家政维修人员。", "系统拒绝提交并提示所选时段无人可接单。", "通过，返回 selected appointment time slot has no available worker。"],
        ["WT05", "P5", "信用分低于阈值的业主提交普通报修。", "系统将普通工单降级为低优先级。", "通过，priority 调整为 3。"],
        ["WT06", "P6", "repair_order.status=处理中，slaDeadline 已超时，30 分钟内无催办记录。", "系统写入 sys_message、repair_urge_log，并将 delayCount 加 1。", "通过，相关记录写入成功。"],
    ]
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    for row_data in new_rows:
        row = table.add_row().cells
        for c, text in enumerate(row_data):
            set_cell_text(row[c], text)

    # Black-box function test table.
    table = doc.tables[12]
    add_rows = [
        ["BT48", "积分与信用记录", "业主完成 5 星维修评价；管理员查询信用分日志。", "系统为该业主增加信用分并写入 credit_log。", "通过"],
        ["BT49", "逾期提醒", "对已逾期账单执行提醒扫描。", "系统生成 payment_reminder 记录，并根据条件生成消息提醒或物业待办。", "通过"],
        ["BT50", "个性化推荐", "业主进入首页，画像满足 hasElderly 或 waterElectricRepairCount 条件。", "系统返回匹配的推荐卡片或弹窗。", "通过"],
        ["BT51", "访客核验", "使用已过期邀请码进行核验。", "系统拒绝通行并返回过期提示。", "通过"],
        ["BT52", "上门核验", "家政维修人员进行活体人脸核验但不匹配。", "系统不写入 verifyPassed，工单不能进入维修处理。", "通过"],
    ]
    for row_data in add_rows:
        row = table.add_row().cells
        for c, text in enumerate(row_data):
            set_cell_text(row[c], text)

    # Equivalence class / boundary / state / scenario / interface tables.
    table = doc.tables[14]
    for row_data in [
        ["EC25", "二手交易发布", "POST /api/neighbor/second-hand/publish，images 为空。", "I1", "提示图片信息缺失或进入无图发布限制流程，具体前端约束需人工确认。", "通过"],
        ["EC26", "访客动态令牌核验", "token 为空。", "V5", "提示 token is required。", "通过"],
        ["EC27", "访客动态令牌核验", "token 为失效或已刷新令牌。", "V6", "提示动态二维码已失效或已刷新。", "通过"],
    ]:
        row = table.add_row().cells
        for c, text in enumerate(row_data):
            set_cell_text(row[c], text)

    table = doc.tables[15]
    for row_data in [
        ["BT22A", "逾期提醒阈值", "未缴金额=1000.00，逾期 30 天。", "系统仍按消息提醒处理，不生成物业待办。", "通过"],
        ["BT22B", "逾期提醒阈值", "未缴金额=1000.01 或逾期 91 天。", "系统生成物业待办。", "通过"],
    ]:
        row = table.add_row().cells
        for c, text in enumerate(row_data):
            set_cell_text(row[c], text)

    table = doc.tables[16]
    row = table.add_row().cells
    for c, text in enumerate(["BT31A", "二手交易审核状态", "已发布", "图片审核结果达到人工复核阈值。", "待审核", "通过"]):
        set_cell_text(row[c], text)

    table = doc.tables[17]
    for row_data in [
        ["BT37A", "积分支付与账单完成", "查询账单→积分支付→扣减积分→更新账单状态。", "points_consumption_record 与 fee_bill 同步更新。", "通过"],
        ["BT37B", "二手交易发布审核", "发布商品→图片检测→返回风险等级→必要时转待审核。", "高风险商品进入人工复核流程。", "通过"],
    ]:
        row = table.add_row().cells
        for c, text in enumerate(row_data):
            set_cell_text(row[c], text)

    table = doc.tables[18]
    set_cell_text(table.cell(0, 1), "接口功能")
    set_cell_text(table.cell(0, 2), "测试情况")
    set_cell_text(table.cell(0, 3), "预期结果")
    set_cell_text(table.cell(1, 1), "二手交易发布接口中的图片审核调用")
    set_cell_text(table.cell(1, 2), "POST /api/neighbor/second-hand/publish，images 为可访问图片地址列表")
    set_cell_text(table.cell(1, 3), "返回 imageAuditResults、imageAuditStatus 和 imageRiskLevel")
    set_cell_text(table.cell(2, 1), "二手交易发布接口中的图片审核调用")
    set_cell_text(table.cell(2, 2), "POST /api/neighbor/second-hand/publish，images 中包含无法读取的图片地址")
    set_cell_text(table.cell(2, 3), "系统把对应图片结果标记为需人工复核")
    set_cell_text(table.cell(3, 1), "模型服务异常容错")
    set_cell_text(table.cell(3, 2), "图片检测服务不可用或返回异常")
    set_cell_text(table.cell(3, 3), "detectLabel=DETECT_FAILED，auditStatus=MANUAL_REVIEW")

    table = doc.tables[19]
    set_cell_text(table.cell(1, 2), "POST /api/auth/account-login，请求体为 account=YZ010110126，password=ZYX123456，role=1。")
    set_cell_text(table.cell(2, 2), "POST /api/repair/submit，请求体缺少 description，其他字段为 serviceMajor=水电维修、appointmentDate=2026-05-18。")
    set_cell_text(table.cell(3, 2), "POST /api/repair/assign，普通业主 Token 提交 orderId=5001，assignee=20。")
    set_cell_text(table.cell(4, 2), "POST /api/fee/pay/wechat，业主重复支付已缴账单 billId=3001。")
    set_cell_text(table.cell(5, 2), "POST /api/access/verify-invite，提交已过期邀请码 code=INV-C001。")
    set_cell_text(table.cell(6, 2), "GET /api/recommend/services，测试用户无匹配画像标签或后台无启用规则。")
    for row_data in [
        ["BT43", "验证码核验接口", "POST /api/worker/verify-code，提交错误的 orderId 与 code。", "返回核验失败，repair_order_worker.verifyPassed 不更新。", "通过"],
        ["BT44", "人脸活体核验接口", "POST /api/face/live-verify-multi，workerId 对应人脸不匹配。", "返回 match=false，不能据此进入维修处理。", "通过"],
        ["BT45A", "SLA 扫描接口", "POST /api/smart-work-orders/sla/scan，工单已超过 slaDeadline。", "返回 delayCount，并生成超时提醒记录。", "通过"],
    ]:
        row = table.add_row().cells
        for c, text in enumerate(row_data):
            set_cell_text(row[c], text)

    # Image audit tables.
    table = doc.tables[22]
    set_cell_text(table.cell(4, 0), "补充伪造样本子集命中率")

    # 4.4 inserted engineering rules.
    anchor = paras[477]
    current = insert_paragraph_after(anchor, "在上述模块类图基础上，结合源码实现，可进一步把 4.4 节中的核心工程规则细化如下。", "Normal")

    current = insert_paragraph_after(current, "（1）智能派单实现规则。业主提交报修后，RepairController 在工单保存阶段调用 pickAssignee 和 collectScoredWorkers 生成建议派工人员；SmartWorkOrderService.dispatch 与 /api/repair/assign 则负责把管理员最终确认的派工结果写回数据库。当前源码先筛选 role=3 且 status=1 的家政维修人员，再读取 worker_staffing.current_status=1 的人员信息，排除当前任务量已达到 max_daily_orders 的候选人；随后根据 staffType、specialties、certificates、是否需要外协和当前未完成工单数进行排序。管理员可采纳 suggested_worker_id，也可手动选择 assigneeIds。派工完成后，repair_order.assignee、repair_order.suggested_worker_id、repair_order.assigned_time、repair_order.status 以及 repair_order_worker 参与人员记录会同步更新。源码未发现历史评分参与推荐，排班 shift_group 也未发现实际参与计算。", "Normal")
    current = insert_paragraph_after(current, "表4-X 维修人员推荐与派工依据表", "Caption")
    current = insert_paragraph_after(current, "Tab.4-X Recommendation and Dispatch Factors for Maintenance Workers", "Caption")
    _, current = insert_table_after(doc, current, [
        ["计算因素", "源码依据", "处理方式"],
        ["专业类型匹配", "worker_staffing.staff_type 与工单服务类型", "匹配时提高得分，不匹配时降低得分"],
        ["专业标签匹配", "worker_staffing.specialties 与 service_sub_type、service_major、category", "子类、主类、分类命中分别提高得分"],
        ["证书匹配", "worker_staffing.certificates 与必需证书关键字", "命中加分，缺失减分"],
        ["当前任务量", "repair_order + repair_order_worker 未完成工单数量", "任务量越少排序越靠前；达到 max_daily_orders 后不再推荐"],
        ["可服务状态", "user.status 与 worker_staffing.current_status", "用户停用或 current_status!=1 时不进入候选集"],
    ])

    current = insert_paragraph_after(current, "（2）工单优先级实现规则。源码中的自动优先级计算由 RepairGradingService.calculatePriority 完成，并在 RepairController.submit 中写入 repair_order.priority。该规则首先对报修描述去除 HTML、标点和空白，再根据紧急关键词、普通关键词和低优先级关键词匹配基础等级；之后结合 user.credit_score 做有限修正：基础为低优先级且信用分不低于 180 时，上调为普通；基础为普通且信用分低于 60 时，下调为低。当前源码未发现预约时间和 SLA 截止时间直接参与自动优先级计算，预约时间主要用于可预约时段校验，SLA 截止时间主要用于后续时效监控。物业管理员可通过 /api/smart-work-orders/{id}/priority 手动调整 priority、suggested_worker_id 和 sla_deadline。", "Normal")
    current = insert_paragraph_after(current, "表4-X 工单优先级判断规则表", "Caption")
    current = insert_paragraph_after(current, "Tab.4-X Priority Judgment Rules for Repair Work Orders", "Caption")
    _, current = insert_table_after(doc, current, [
        ["判断项", "源码依据", "处理结果"],
        ["紧急关键词", "漏水、爆管、断电、漏电、门锁故障等关键字集合", "priority=1"],
        ["普通维修关键词", "灯具、水龙头、开关、插座、堵塞等", "priority=2"],
        ["咨询或预约类关键词", "咨询、报价、预约、上门时间等", "priority=3"],
        ["高信用用户", "基础等级为低且 credit_score>=180", "上调为普通优先级"],
        ["低信用用户", "基础等级为普通且 credit_score<60", "下调为低优先级"],
        ["预约时间与 SLA", "未直接进入自动优先级计算", "由可预约时段校验和时效监控单独处理"],
    ])

    current = insert_paragraph_after(current, "（3）积分与信用分记录规则。当前源码将“积分”和“信用分”分开实现。user.points、points_recharge_record 和 points_consumption_record 用于积分充值与积分消费；user.credit_score、credit_log 和 CreditService 用于行为信用记录。已能从源码确认的信用分变动场景包括：报修五星评价后加 2 分、二手交易标记已售后加 5 分、物业管理员人工调整，以及按月执行的信用分衰减。当前未发现缴费、投诉、失物招领直接触发信用分增减的统一实现，该部分需人工确认。", "Normal")
    current = insert_paragraph_after(current, "表4-X 积分与信用分处理规则表", "Caption")
    current = insert_paragraph_after(current, "Tab.4-X Points and Credit Processing Rules", "Caption")
    _, current = insert_table_after(doc, current, [
        ["行为类型", "对应服务或表", "处理结果"],
        ["物业管理员积分充值", "PointsService.recharge / points_recharge_record", "增加 user.points，并记录前后余额和备注"],
        ["账单、停车费或维修费积分支付", "PointsService.consume / points_consumption_record", "扣减 user.points，并记录 businessType、businessId 和前后余额"],
        ["报修五星评价", "RepairEvaluationServiceImpl + CreditService", "增加 2 分信用分，并写入 credit_log.reason"],
        ["二手交易成功", "NeighborServiceImpl + CreditService", "增加 5 分信用分，并写入 credit_log.reason"],
        ["管理员人工调整", "CreditController.adjust", "按请求增减信用分，并写入变动原因"],
        ["信用分时间衰减", "CreditDecayTask + UserMapper.decayCreditScores", "按 0.98^(天数/30) 近似衰减，范围限制在 0~200"],
    ])

    current = insert_paragraph_after(current, "（4）逾期提醒实现规则。PaymentReminderService 通过 @Scheduled(cron = \"0 0 8 * * ?\") 每天 8:00 扫描逾期账单。服务先调用 PaymentReminderMapper.selectOverdueFeeBillsForReminder 查询未缴且已逾期的 fee_bill，再计算未缴金额和逾期天数。当前源码实现为两档提醒：未缴金额大于 1000 元或逾期超过 90 天时生成 TASK 类型提醒并创建 property_task；其余存在欠费的逾期账单生成 MESSAGE 类型提醒并写入 sys_message。生成前会检查当日是否已存在同类消息提醒、或是否已有未处理的待办提醒，以避免重复创建。payment_reminder 表当前未单独保存催缴等级字段。", "Normal")
    current = insert_paragraph_after(current, "表4-X 逾期提醒处理规则表", "Caption")
    current = insert_paragraph_after(current, "Tab.4-X Overdue Payment Reminder Rules", "Caption")
    _, current = insert_table_after(doc, current, [
        ["判断条件", "提醒方式", "落库结果"],
        ["未缴金额>0，且未达到待办阈值", "MESSAGE", "写入 payment_reminder(method=MESSAGE,status=1) 并生成 sys_message"],
        ["未缴金额>1000 元，或 overdueDays>90", "TASK", "写入 payment_reminder(method=TASK,status=0) 并生成 property_task"],
        ["同一账单同一方式已存在当日消息提醒", "跳过生成", "避免重复消息提醒"],
        ["同一账单已存在未处理 TASK 提醒", "跳过生成", "避免重复创建物业待办"],
    ])

    current = insert_paragraph_after(current, "（5）个性化推荐实现规则。RecommendController 对外提供 /api/recommend/services 和 /api/recommend/popup 两个接口，分别返回首页推荐和弹窗推荐。RecommendRuleService 读取 recommend_rule 表中启用规则，并构造包含 hasElderly、hasChild、hasPet、houseArea、roomCount、水电类报修次数以及隐式关键词的上下文。显式画像来自 user 表字段；水电类报修次数来自 repair_order 统计；隐式关键词由 ImplicitProfileService 每天凌晨从 second_hand.title 和 lost_found.description 中提取，并写入 user_implicit_profile.keywords_json。首页普通推荐默认返回不超过 3 条内容，弹窗推荐默认返回不超过 2 条内容，并按 service_id 去重。当前未发现跨会话的“避免重复打扰”记录机制，该部分需人工确认。", "Normal")
    current = insert_paragraph_after(current, "表4-X 个性化推荐规则示例表", "Caption")
    current = insert_paragraph_after(current, "Tab.4-X Examples of Personalized Recommendation Rules", "Caption")
    _, current = insert_table_after(doc, current, [
        ["画像或条件", "来源", "推荐内容"],
        ["hasElderly=true", "user.has_elderly", "适老化改造服务"],
        ["waterElectricRepairCount>=2", "repair_order 统计结果", "水电维修保养服务"],
        ["houseArea>=120", "user.house_area", "深度保洁服务"],
        ["keywordsAny 命中“宠物”", "user_implicit_profile.keywords_json", "宠物除螨或照看服务"],
    ])

    current = insert_paragraph_after(current, "（6）二手交易图片真实性辅助审核实现规则。NeighborServiceImpl 在发布和修改二手交易信息时调用 ImageAuditService.auditSecondHandImages，对 second_hand.images 中的图片地址逐张执行检测。ImageAuditClient 通过 HTTP Multipart 请求调用模型服务接口，接收 detectLabel、realProbability、fakeProbability 和 threshold；ImageAuditService 再依据 image-audit.fake-threshold=0.5 和 image-audit.review-threshold=0.7 生成风险等级与审核状态，并把结果写入 image_audit_result 表。当前源码内部审核状态为 PASS、SUSPICIOUS 和 MANUAL_REVIEW，对论文表述可统一为正常、疑似异常和需人工复核。若任一图片达到 MANUAL_REVIEW 条件，则 second_hand.status 被调整为待审核。源码未发现独立的二手交易图片表，商品图片地址保存在 second_hand.images 字段中。", "Normal")

    current = insert_paragraph_after(current, "（7）SLA 时效监控实现规则。RepairController.submit 在工单创建时将 sla_deadline 设置为当前时间后 30 分钟，用于待派单阶段；管理员完成派工后，RepairController.assign、RepairController.autoAssign 或 SmartWorkOrderService.dispatch 会把 sla_deadline 更新为当前时间后 2 小时，用于处理中阶段。SlaMonitorService 以每分钟一次的定时任务扫描 repair_order 中 status 为待派单或处理中、且 sla_deadline 已超时的工单。若工单处于待派单状态，则向物业管理员发送超时消息；若工单处于处理中状态，则同时向家政维修人员和物业管理员发送提醒，并写入 repair_urge_log。为避免重复催办，系统要求 30 分钟内不存在同一工单的催办日志后才再次发送提醒。", "Normal")
    current = insert_paragraph_after(current, "表4-X SLA 时效监控规则表", "Caption")
    current = insert_paragraph_after(current, "Tab.4-X SLA Monitoring Rules", "Caption")
    _, current = insert_table_after(doc, current, [
        ["场景", "判定条件", "系统处理"],
        ["待派单时限", "submit 后 30 分钟内未派单", "repair_order.sla_deadline=createTime+30min"],
        ["处理中时限", "派工后 2 小时内未完成", "repair_order.sla_deadline=assignedTime+2h"],
        ["待派单超时", "status=待派单 且当前时间>sla_deadline", "向物业管理员发送 sys_message，写入 repair_urge_log，delay_count+1"],
        ["处理中超时", "status=处理中 且当前时间>sla_deadline", "向家政维修人员和物业管理员发送提醒，写入 repair_urge_log，delay_count+1"],
        ["重复催办抑制", "30 分钟内已存在催办日志", "本轮扫描跳过该工单"],
    ])

    doc.save(str(OUTPUT_PATH))
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
